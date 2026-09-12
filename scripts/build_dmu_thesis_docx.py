"""Build the dissertation in the supplied DMU MSc thesis style."""
from __future__ import annotations

import argparse
import csv
import json
import re
from copy import deepcopy
from html import escape
from pathlib import Path

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASE = ROOT / "docs" / "When_Explanations_Outlive_Their_Data_DMU_Thesis.docx"
DEFAULT_OUTPUT = ROOT / "docs" / "When_Explanations_Outlive_Their_Data_DMU_Thesis_Final.docx"
MATHML_TO_OMML = Path(
    "/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl"
)

CITATIONS = {
    1: "Lin et al., 2018",
    2: "Qin, Zhu and Ye, 2022",
    3: "Lowe et al., 2017",
    4: "Rashid et al., 2020",
    5: "Yu et al., 2022",
    6: "Scarselli et al., 2009",
    7: "Veličković et al., 2018",
    8: "Iqbal and Sha, 2019",
    9: "Vaswani et al., 2017",
    10: "Heuillet, Couthouis and Díaz-Rodríguez, 2021",
    11: "Madumal et al., 2020",
    12: "Jain and Wallace, 2019",
    13: "Wiegreffe and Pinter, 2019",
    14: "Serrano and Smith, 2019",
    15: "Liu et al., 2022",
    16: "DeYoung et al., 2020",
    17: "Jacovi and Goldberg, 2020",
    18: "Ribeiro, Singh and Guestrin, 2016",
    20: "Lipton, 2018",
    21: "Kaul, Yates and Gruteser, 2012",
    22: "Lopez et al., 2018",
    23: "Ying et al., 2019",
    24: "Luo et al., 2020",
    25: "Yuan et al., 2023",
    26: "Abnar and Zuidema, 2020",
    27: "Adebayo et al., 2018",
    28: "Hooker et al., 2019",
    29: "Amara et al., 2022",
    30: "Milani et al., 2024",
    31: "Bekkemoen, 2024",
    33: "Hu, Feng and Li, 2025",
    34: "Greydanus et al., 2018",
    35: "Mott et al., 2019",
    36: "Yuan et al., 2021",
    37: "Wang et al., 2025",
    38: "Sha et al., 2026",
    39: "Yates et al., 2021",
    40: "Bouteiller et al., 2021",
    41: "Liotet et al., 2022",
    42: "Tabassi, 2023",
    43: "Zheng et al., 2024",
    44: "Li et al., 2024",
    45: "Azzolin et al., 2025",
    46: "Shin et al., 2025",
    47: "Azzolin et al., 2026",
    48: "Ancona et al., 2018",
    49: "S. Lu et al., 2024",
    50: "Bui et al., 2024",
    51: "Zhang et al., 2024",
    52: "Armgaan et al., 2024",
    53: "Liu and Xie, 2025",
    54: "Liu et al., 2025",
    55: "Saha and Bandyopadhyay, 2026",
    56: "Wu et al., 2025",
    57: "Zhou et al., 2024",
    58: "Yao, Florescu and Lee, 2024",
    59: "Ding et al., 2024",
    60: "Soudijani and Dimitrova, 2025",
    61: "Hong et al., 2024",
    62: "W. Lu et al., 2024",
    63: "Amitai, Septon and Amir, 2024",
    64: "Pan et al., 2025",
    65: "He et al., 2025",
    66: "Varbella et al., 2024",
    67: "Wang and Shen, 2024",
    68: "Yu and Gao, 2025",
    69: "OpenStreetMap contributors, n.d.",

}

TABLE_TITLES = {
    (2, 1): "Critical comparison of eight closely related explanation studies",
    (3, 1): "Observation graph node types and roles",
    (3, 2): "Policy conditions used in the experiment",
    (3, 3): "Validation-selected checkpoint indices (zero-based)",
    (3, 4): "Training and optimization settings",
    (3, 5): "Checkpoint-acceptance criteria",
    (3, 6): "Explanation measures and roles",
    (3, 7): "Held-out evaluation structure and episode counts",
    (3, 8): "Telemetry conditions used in held-out evaluation",
    (3, 9): "Hypothesis map and analysis roles",
    (3, 10): "Implementation components and design rationale",
    (4, 1): "Policy capability on held-out demand",
    (4, 2): "Protocol correction and numeric-precision audit",
    (4, 3): "Decision-relevance controls under clean and degraded telemetry",
    (4, 4): "Valid node-count distribution in scored graphs",
    (4, 5): "Exposure-conditioned paired attention and DEF shifts",
    (4, 6): "Action composition and paired DEF by action type",
    (4, 7): "Tunnel and random-trigger sensitivity summary",
    (4, 8): "Hypothesis outcomes across training seeds",
    (4, 9): "Freshness-aware explanation audit decisions",
    (4, 10): "Exploratory no-op margin-DEF by checkpoint and condition",
    (5, 1): "Freshness-aware explanation audit framework",
    (5, 2): "Audit framework inputs and outputs",
}

FIGURES = [
    ("Figure 1.1", "Stale data and the explanation-release decision"),
    ("Figure 3.1", "Spatial distribution of generated demand and tunnel triggers"),
    ("Figure 3.2", "Demand arrival and origin-destination distributions by split"),
    ("Figure 3.3", "Local observation graph and request-to-action mapping"),
    ("Figure 3.4", "Graph-attention policy and audit channel"),
    ("Figure 3.5", "Training, selection and held-out evaluation"),
    ("Figure 3.6", "Tunnel-triggered observation-layer telemetry degradation"),
    ("Figure 3.7", "Construct-validity controls for action-linked request nodes"),
    ("Figure 4.1", "GAT and GAT-Outage training and checkpoint-selection diagnostics"),
    ("Figure 4.2", "Clean-telemetry completed journeys by training seed"),
    ("Figure 4.3", "Faithfulness perturbation controls by checkpoint"),
    ("Figure 4.4", "Attention aggregation sensitivity"),
    ("Figure 4.5", "Attention query-row sensitivity"),
    ("Figure 4.6", "Top-k overlap and DEF resolution"),
    ("Figure 4.7", "Outage-duration sweep by policy and training seed"),
    ("Figure 4.8", "Paired attention and faithfulness shift by training seed"),
    ("Figure 4.9", "Faithfulness analysis by action type"),
    ("Figure 4.10", "Tunnel-triggered and random telemetry-loss sensitivity"),
    ("Figure 4.11", "WAMSN-DEF relationship and H4 result"),
    ("Figure 5.1", "Evidence interpretation summary"),
    ("Figure 5.2", "Freshness-aware explanation audit workflow"),
]

TABLES = [(f"Table {chapter}.{number}", title)
          for (chapter, number), title in TABLE_TITLES.items()]
TABLES.append(("Table C.1", "Rerun per-checkpoint hypothesis statistics"))

FIGURE_PAGES = {
    "Figure 1.1": "3",
    "Figure 3.1": "11", "Figure 3.2": "13", "Figure 3.3": "15",
    "Figure 3.4": "16", "Figure 3.5": "19",
    "Figure 4.1": "25", "Figure 4.2": "26", "Figure 4.3": "28",
    "Figure 4.4": "29", "Figure 4.5": "30",
    "Figure 4.6": "31", "Figure 4.7": "32", "Figure 4.8": "33",
    "Figure 4.9": "34", "Figure 4.10": "35", "Figure 4.11": "35",
    "Figure 5.1": "39", "Figure 5.2": "41",
}

TABLE_PAGES = {
    "Table 2.1": "9", "Table 3.1": "11", "Table 3.2": "13",
    "Table 3.3": "14", "Table 3.4": "14", "Table 3.5": "14",
    "Table 3.6": "18", "Table 3.7": "20", "Table 3.8": "20",
    "Table 3.9": "21",
    "Table 4.1": "26", "Table 4.2": "27",
    "Table 4.3": "27", "Table 4.4": "30", "Table 4.5": "33",
    "Table 4.6": "34", "Table 4.7": "35", "Table 4.8": "36",
    "Table 4.9": "37", "Table 5.1": "41", "Table 5.2": "42",
}

HEADING_PAGES = {
    "Chapter 1: Introduction": "1", "1.1 Context": "1",
    "1.2 Problem statement": "1", "1.3 Research questions": "1",
    "1.4 Objectives": "2", "1.5 Main findings": "2",
    "1.6 Contributions": "3", "1.7 Dissertation structure": "4",
    "Chapter 2: Literature Review": "5",
    "2.1 Reinforcement learning for fleet dispatch": "5",
    "2.2 Explainability in reinforcement learning": "5",
    "2.3 Is attention an explanation?": "6",
    "2.4 Explaining graph neural networks": "7",
    "2.5 Evaluating faithfulness: methods and pitfalls": "7",
    "2.6 Age of Information and telemetry degradation": "8",
    "2.7 From explanation evaluation to explanation assurance": "8",
    "2.8 Research gap and positioning": "9",
    "Chapter 3: Methodology": "10", "3.1 Study design": "10",
    "3.2 SUMO environment and data": "10",
    "3.3 Observation graph and action space": "10",
    "3.4 Policy models and training": "11", "3.5 Telemetry degradation": "15",
    "3.6 Explanation measures": "16",
    "3.6.1 Decision-level explanation faithfulness": "16",
    "3.6.2 Stale-node attention": "18", "3.7 Construct-validity audit": "18",
    "3.8 Evaluation design": "19", "3.8.1 Primary evaluation": "19",
    "3.8.2 Supporting analyses": "21",
    "3.9 Hypotheses and statistical analysis": "21",
    "3.9.1 Statistical procedure": "22", "3.9.2 Interpretation limits": "22",
    "3.10 Audit decision implementation": "22", "3.11 Reproducibility": "23",
    "3.12 Ethics and data governance": "24",
    "Chapter 4: Results": "25",
    "4.1 Policy capability and training stability": "25",
    "4.2 Protocol correction and precision audit": "26",
    "4.3 Decision-relevance controls": "27",
    "4.4 Small-graph resolution": "30",
    "4.5 Telemetry manipulation and stale exposure": "31",
    "4.6 Paired clean and degraded observations": "32",
    "4.7 Analysis by action type": "33",
    "4.8 Random-trigger sensitivity analysis": "34",
    "4.9 Hypothesis results": "35", "4.10 Answers to the research questions": "36",
    "4.11 Explanation-release audit": "37",
    "Chapter 5: Discussion": "38", "5.1 Answer to the central problem": "38",
    "5.2 How to read the evidence": "38",
    "5.3 Why the result is not consistent": "39",
    "5.4 Why the construct-validity controls matter": "40",
    "5.5 Why outage training did not solve the problem": "40",
    "5.6 Proposed audit framework": "40",
    "5.6.1 Purpose and scope": "40", "5.6.2 Inputs, outputs and use": "42",
    "5.6.3 Application to this study": "42",
    "5.7 Limitations": "43", "5.8 Future work": "43",
    "Chapter 6: Conclusion": "45",
}


def paragraph_text(element) -> str:
    # Word may repeat text in proofing or field XML; Paragraph.text reflects
    # the single displayed string.
    return Paragraph(element, None).text.strip()


def find_body_child(doc: Document, text: str):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p") and paragraph_text(child) == text:
            return child
    raise ValueError(f"paragraph not found: {text}")


def clear_between(doc: Document, start_text: str, end_text: str) -> tuple[object, object]:
    body = doc.element.body
    start = find_body_child(doc, start_text)
    end = find_body_child(doc, end_text)
    children = list(body.iterchildren())
    start_index = children.index(start)
    end_index = children.index(end)
    for child in children[start_index + 1:end_index]:
        body.remove(child)
    return start, end


def remove_range_inclusive(doc: Document, start_text: str, end_text: str) -> object:
    body = doc.element.body
    start = find_body_child(doc, start_text)
    end = find_body_child(doc, end_text)
    children = list(body.iterchildren())
    start_index = children.index(start)
    end_index = children.index(end)
    for child in children[start_index:end_index]:
        body.remove(child)
    return end


def move_before(element, anchor) -> None:
    anchor.addprevious(element)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, *, top: int = 45, start: int = 90,
                     bottom: int = 45, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    settings = {
        "top": ("single", "8", "7F7F7F"),
        "bottom": ("single", "8", "7F7F7F"),
        "left": ("nil", "0", "FFFFFF"),
        "right": ("nil", "0", "FFFFFF"),
        "insideH": ("single", "4", "D9D9D9"),
        "insideV": ("nil", "0", "FFFFFF"),
    }
    for side, (style, size, colour) in settings.items():
        border = borders.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), style)
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), colour)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{side}"))
        if border is None:
            border = OxmlElement(f"w:{side}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def _mathml_tag(name: str, content: str) -> str:
    return f"<{name}>{content}</{name}>"


def _mi(value: str) -> str:
    return _mathml_tag("mi", escape(value))


def _mn(value: str | int) -> str:
    return _mathml_tag("mn", escape(str(value)))


def _mo(value: str) -> str:
    return _mathml_tag("mo", escape(value))


def _mtext(value: str) -> str:
    return _mathml_tag("mtext", escape(value))


def _mrow(*parts: str) -> str:
    return _mathml_tag("mrow", "".join(parts))


def _sub(base: str, subscript: str) -> str:
    return _mathml_tag("msub", base + subscript)


def _sup(base: str, superscript: str) -> str:
    return _mathml_tag("msup", base + superscript)


def _subsup(base: str, subscript: str, superscript: str) -> str:
    return _mathml_tag("msubsup", base + subscript + superscript)


def _frac(numerator: str, denominator: str) -> str:
    return _mathml_tag("mfrac", numerator + denominator)


def _sqrt(value: str) -> str:
    return _mathml_tag("msqrt", value)


def _call(name: str, *arguments: str) -> str:
    joined = []
    for index, argument in enumerate(arguments):
        if index:
            joined.append(_mo(","))
        joined.append(argument)
    return _mrow(_mtext(name), _mo("("), *joined, _mo(")"))


def _matrix(*rows: str) -> str:
    body = "".join(
        _mathml_tag("mtr", _mathml_tag("mtd", row)) for row in rows
    )
    return f'<mtable columnalign="left">{body}</mtable>'


def _mathml(body: str) -> str:
    return (
        '<math xmlns="http://www.w3.org/1998/Math/MathML" display="block">'
        + body
        + "</math>"
    )


def equation_spec(code: str) -> tuple[str, str] | None:
    """Map methodology formula blocks to numbered native Word equations."""
    comma_h = _mrow(_mi("Q"), _mo(","), _mi("h"))
    h_i_l = _subsup(_mi("h"), _mi("i"), _mi("l"))
    h_j_l = _subsup(_mi("h"), _mi("j"), _mi("l"))
    q_i = _sub(_mi("q"), _mi("i"))
    k_j = _sub(_mi("k"), _mi("j"))
    v_j = _sub(_mi("v"), _mi("j"))
    alpha = _subsup(
        _mi("α"), _mrow(_mi("i"), _mi("j")),
        _mrow(_mi("l"), _mo(","), _mi("h")),
    )

    if code.startswith("q_i ="):
        rows = (
            _mrow(q_i, _mo("="), _sub(_mi("W"), comma_h), _call("LayerNorm", h_i_l)),
            _mrow(k_j, _mo("="), _sub(_mi("W"), _mrow(_mi("K"), _mo(","), _mi("h"))), _call("LayerNorm", h_j_l)),
            _mrow(v_j, _mo("="), _sub(_mi("W"), _mrow(_mi("V"), _mo(","), _mi("h"))), _call("LayerNorm", h_j_l)),
            _mrow(
                _subsup(_mi("s"), _mrow(_mi("i"), _mi("j")), _mrow(_mi("l"), _mo(","), _mi("h"))),
                _mo("="),
                _frac(_mrow(q_i, k_j), _sqrt(_sub(_mi("d"), _mtext("head")))),
            ),
            _mrow(alpha, _mo("="), _sub(_mtext("softmax"), _mi("j")), _mo("("), _subsup(_mi("s"), _mrow(_mi("i"), _mi("j")), _mrow(_mi("l"), _mo(","), _mi("h"))), _mo(")")),
            _mrow(
                _sub(_mi("z"), _mi("i")), _mo("="),
                _sub(_mtext("concat"), _mi("h")), _mo("("),
                _sub(_mi("Σ"), _mi("j")), _mo("("), alpha, _mo("×"), v_j,
                _mo(")"), _mo(")"),
            ),
            _mrow(_sub(_mi("u"), _mi("i")), _mo("="), h_i_l, _mo("+"), _sub(_mi("W"), _mi("O")), _sub(_mi("z"), _mi("i"))),
            _mrow(_subsup(_mi("h"), _mi("i"), _mrow(_mi("l"), _mo("+"), _mn(1))), _mo("="), _sub(_mi("u"), _mi("i")), _mo("+"), _call("FFN", _call("LayerNorm", _sub(_mi("u"), _mi("i"))))),
        )
        return "(3.1)", _mathml(_matrix(*rows))

    if code.startswith("alpha_j ="):
        mean = _frac(
            _mrow(
                _sub(_mi("Σ"), _mrow(_mi("l"), _mo(","), _mi("h"))),
                _mo("("),
                _subsup(_mi("α"), _mrow(_mn(0), _mi("j")), _mrow(_mi("l"), _mo(","), _mi("h"))),
                _mo(")"),
            ),
            _mrow(_mi("L"), _mi("H")),
        )
        return "(3.2)", _mathml(_mrow(_sub(_mi("α"), _mi("j")), _mo("="), _call("normalize", mean)))

    if code.startswith("r_t ="):
        formula = _mrow(
            _sub(_mi("r"), _mi("t")), _mo("="), _mn(10), _sub(_mi("N"), _mrow(_mtext("complete"), _mo(","), _mi("t"))),
            _mo("+"), _mn("0.5"), _sub(_mi("N"), _mrow(_mtext("dispatch"), _mo(","), _mi("t"))),
            _mo("-"), _mn("0.001"), _sub(_mtext("mean_wait"), _mi("t")),
        )
        return "(3.3)", _mathml(formula)

    if code.startswith("r_train,i,t ="):
        formula = _mrow(
            _subsup(
                _mi("r"),
                _mrow(_mi("i"), _mo(","), _mi("t")),
                _mtext("train"),
            ),
            _mo("="),
            _sub(_mi("r"), _mi("t")),
            _mo("+"),
            _mn("1.0"),
            _sub(_mi("I"), _mtext("successful dispatch by i at t")),
        )
        return "(3.4)", _mathml(formula)

    if code.startswith("Comp(R_k; a)"):
        r_k = _sub(_mi("R"), _mi("k"))
        b_k_b = _subsup(_mi("B"), _mi("k"), _mi("b"))
        action = _mi("a")
        score = lambda graph: _call("s", graph, _mi("a"))
        rows = (
            _mrow(_call("Comp", r_k, action), _mo("="), score(_mi("G")), _mo("-"), score(_mrow(_mi("G"), _mo("∖"), r_k))),
            _mrow(_call("Suff", r_k, action), _mo("="), score(_mi("G")), _mo("-"), score(_mrow(_mi("G"), _mo("["), r_k, _mo("]")))),
            _mrow(_sub(_mi("g"), _mtext("comp")), _mo("("), _mi("k"), _mo(";"), action, _mo(")"), _mo("="), _call("Comp", r_k, action), _mo("-"), _sub(_mtext("mean"), _mi("b")), _call("Comp", b_k_b, action)),
            _mrow(_sub(_mi("g"), _mtext("suff")), _mo("("), _mi("k"), _mo(";"), action, _mo(")"), _mo("="), _sub(_mtext("mean"), _mi("b")), _call("Suff", b_k_b, action), _mo("-"), _call("Suff", r_k, action)),
            _mrow(_call("DEF", action), _mo("="), _sub(_mtext("mean"), _mi("k")), _frac(_mn(1), _mn(2)), _mo("["), _sub(_mi("g"), _mtext("comp")), _mo("("), _mi("k"), _mo(";"), action, _mo(")"), _mo("+"), _sub(_mi("g"), _mtext("suff")), _mo("("), _mi("k"), _mo(";"), action, _mo(")"), _mo("]")),
        )
        return "(3.5)", _mathml(_matrix(*rows))

    if code.startswith("WAMSN ="):
        numerator = _mrow(
            _sub(_mi("Σ"), _mi("i")), _mo("("),
            _sub(_mtext("attention"), _mi("i")), _mo("×"),
            _sub(_mtext("normalized_AoI"), _mi("i")), _mo(")"),
        )
        denominator = _mrow(
            _sub(_mi("Σ"), _mi("i")), _mo("("),
            _sub(_mtext("attention"), _mi("i")), _mo(")"),
        )
        return "(3.6)", _mathml(_mrow(_mtext("WAMSN"), _mo("="), _frac(_mrow(numerator), _mrow(denominator))))

    count_specs = {
        "3 models x": ("(3.7)", "3 x 5 x 8 x 6 = 720"),
        "2 models x 5 training seeds x 5 conditions": ("(3.8)", "2 x 5 x 5 x 8 x 6 = 2,400"),
        "2 models x 5 training seeds x 2 conditions": ("(3.9)", "2 x 5 x 2 x 3 x 3 = 180"),
    }
    for prefix, (number, expression) in count_specs.items():
        if code.startswith(prefix):
            parts = expression.replace(" x ", " × ").split(" = ")
            return number, _mathml(_mrow(_mtext(parts[0]), _mo("="), _mn(parts[1])))
    return None


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
    set_table_borders(table)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def keep_row(row) -> None:
    prevent_row_split(row)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_together = True


def add_numbering_definition(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(element.get(qn("w:abstractNumId")))
                    for element in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(element.get(qn("w:numId")))
               for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    paragraph_properties.append(indent)
    level.append(paragraph_properties)
    if bullet:
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "Times New Roman")
        fonts.set(qn("w:cs"), "Times New Roman")
        run_properties.append(fonts)
        level.append(run_properties)
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id


def ensure_list_style(doc: Document, name: str, *, bullet: bool):
    styles = doc.styles
    if name in [style.name for style in styles]:
        style = styles[name]
    else:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    run_properties = style._element.get_or_add_rPr()
    language = run_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        run_properties.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "en-US")
    language.set(qn("w:bidi"), "en-US")
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.widow_control = True

    p_pr = style._element.get_or_add_pPr()
    existing_numbering = p_pr.find(qn("w:numPr"))
    if existing_numbering is not None:
        p_pr.remove(existing_numbering)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(add_numbering_definition(doc, bullet=bullet)))
    num_pr.append(num_id)
    p_pr.append(num_pr)
    return style


def set_paragraph_numbering(paragraph, num_id: int) -> None:
    """Assign a list instance so each independent numbered list restarts at 1."""
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_pr.append(level)
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(number)
    p_pr.append(num_pr)


def apply_inline(paragraph, text: str) -> None:
    """Add a small Markdown subset while preserving DMU body typography."""
    text = text.replace("`", "")
    pattern = re.compile(r"(https?://\S+|\*\*.+?\*\*|\*.+?\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith(("https://", "http://")):
            url = token.rstrip(".,;")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = paragraph.add_run(url)
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
            if len(url) < len(token):
                paragraph.add_run(token[len(url):])
            position = match.end()
            continue
        run = paragraph.add_run(token.strip("*"))
        run.bold = token.startswith("**")
        run.italic = not token.startswith("**")
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def convert_citations(text: str) -> str:
    # Preserve narrative Harvard citations: "Author [n]" becomes
    # "Author (year)", while a standalone [n] remains "(Author, year)".
    ordered_groups = ([6, 7], [12, 14, 13, 15], [12, 14, 15], [12, 13])
    for group in ordered_groups:
        token = ", ".join(f"[{number}]" for number in group)
        citation = "; ".join(CITATIONS[number] for number in group)
        text = text.replace(token, f"({citation})")

    for number, citation in CITATIONS.items():
        authors, year = citation.rsplit(", ", 1)
        text = text.replace(f"{authors} [{number}]", f"{authors} ({year})")

    def replace_range(match: re.Match[str]) -> str:
        start, end = int(match.group(1)), int(match.group(2))
        return "(" + "; ".join(CITATIONS[number] for number in range(start, end + 1)) + ")"

    text = re.sub(r"\[(\d+)\]\s*[-–]\s*\[(\d+)\]", replace_range, text)

    def replace_single(match: re.Match[str]) -> str:
        number = int(match.group(1))
        return f"({CITATIONS[number]})" if number in CITATIONS else match.group(0)

    return re.sub(r"\[(\d+)\]", replace_single, text)


def style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for name, size, alignment in (
        ("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.space_after = Pt(12)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].paragraph_format.space_before = Pt(9)
    styles["Heading 3"].paragraph_format.space_after = Pt(3)

    ensure_list_style(doc, "DMU Bullet", bullet=True)
    ensure_list_style(doc, "DMU Number", bullet=False)

    if "DMU Caption" not in [style.name for style in styles]:
        styles.add_style("DMU Caption", 1)
    caption = styles["DMU Caption"]
    caption.font.name = "Times New Roman"
    caption.font.size = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(0.59)
        section.bottom_margin = Inches(1.03)


def add_paragraph_before(doc: Document, anchor, text: str = "", style: str = "Normal"):
    paragraph = doc.add_paragraph(style=style)
    if text:
        apply_inline(paragraph, text)
    move_before(paragraph._p, anchor)
    return paragraph


def add_caption_before(doc: Document, anchor, text: str):
    paragraph = add_paragraph_before(doc, anchor, style="DMU Caption")
    paragraph.add_run(text).bold = True
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_table_before(doc: Document, anchor, rows: list[list[str]], title: str):
    caption_paragraph = add_caption_before(doc, anchor, title)
    if title.startswith(("Table 2.1:", "Table 4.8:", "Table 4.10:")):
        caption_paragraph.paragraph_format.page_break_before = True
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = None
    total = 9020
    if title.startswith("Table 2.1:"):
        widths = [1800, 1700, 2000, 1400, 2120]
    elif title.startswith("Table 3.2:"):
        widths = [1200, 1650, 2650, 3520]
    elif title.startswith("Table 3.5:"):
        widths = [5900, 3120]
    elif title.startswith("Table 3.6:"):
        widths = [2100, 2850, 4070]
    elif title.startswith("Table 3.7:"):
        widths = [2800, 1000, 1400, 1100, 1500, 1220]
    elif title.startswith("Table 3.8:"):
        widths = [1100, 2500, 2000, 3420]
    elif title.startswith("Table 3.9:"):
        widths = [1100, 2900, 3000, 2020]
    elif title.startswith("Table 4.6:"):
        widths = [1600, 850, 850, 1900, 1900, 1920]
    elif title.startswith("Table 4.9:"):
        widths = [1700, 1850, 1850, 3620]
    elif title.startswith("Table 5.1:"):
        widths = [1900, 2350, 1900, 2870]
    else:
        weights = []
        for column in range(len(rows[0])):
            longest = max(len(row[column]) for row in rows)
            weights.append(max(8, min(longest, 45)))
        widths = [int(total * weight / sum(weights)) for weight in weights]
        widths[-1] += total - sum(widths)
    set_table_geometry(table, widths)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        keep_row(row)
        if row_index == 0:
            repeat_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.widow_control = True
            center_value = column_index > 0 and (
                title.startswith("Table 4.6:") or len(value) < 25
            )
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if center_value else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(convert_citations(value))
            run.font.name = "Times New Roman"
            compact_table = title.startswith(("Table 3.9:", "Table 4.9:"))
            run.font.size = Pt(9 if compact_table else 10)
            if title.startswith("Table 4.9:") and value == "INDETERMINATE":
                run.font.size = Pt(8)
            if row_index == 0:
                run.bold = True
                set_cell_shading(cell, "E7E6E6")
    keep_complete = title.startswith(("Table 4.3:", "Table 4.5:", "Table 4.6:"))
    if len(rows) <= 4 or keep_complete:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    move_before(table._tbl, anchor)
    return caption_paragraph


def add_equation_before(doc: Document, anchor, code: str) -> bool:
    spec = equation_spec(code)
    if spec is None:
        return False
    if not MATHML_TO_OMML.exists():
        raise FileNotFoundError(
            "Microsoft Word MathML-to-OMML stylesheet was not found: "
            f"{MATHML_TO_OMML}"
        )

    number, mathml = spec
    transform = etree.XSLT(etree.parse(str(MATHML_TO_OMML)))
    omml = transform(etree.fromstring(mathml.encode("utf-8"))).getroot()

    table = doc.add_table(rows=1, cols=3)
    table.style = None
    set_table_geometry(table, [650, 7720, 650])
    remove_table_borders(table)
    row = table.rows[0]
    prevent_row_split(row)
    for cell in row.cells:
        set_cell_margins(cell, top=30, start=0, bottom=30, end=0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.keep_together = True

    equation_paragraph = row.cells[1].paragraphs[0]
    equation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation_paragraph._p.append(deepcopy(omml))

    number_paragraph = row.cells[2].paragraphs[0]
    number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_run = number_paragraph.add_run(number)
    number_run.font.name = "Times New Roman"
    number_run.font.size = Pt(11)

    move_before(table._tbl, anchor)
    return True


def parse_table(lines: list[str]) -> list[list[str]]:
    parsed = [[cell.strip() for cell in line.strip().strip("|").split("|")]
              for line in lines]
    return [parsed[0], *parsed[2:]]


def add_image_before(doc: Document, anchor, source: Path, alt_text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    width = 5.5 if source.name.endswith("gat_training_diagnostics.png") else 6.15
    run.add_picture(str(source), width=Inches(width))
    drawing = run._r.find(qn("w:drawing"))
    if drawing is not None:
        doc_pr = drawing.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", alt_text)
    move_before(paragraph._p, anchor)
    return paragraph


def collect_blocks(path: Path) -> list[tuple[str, object]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[tuple[str, object]] = []
    index = 0
    paragraph_lines: list[str] = []

    def flush() -> None:
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines)
            blocks.append(("paragraph", convert_citations(text)))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith("```"):
            flush()
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(("code", "\n".join(code)))
            index += 1
            continue
        image_match = re.match(r"!\[(.+?)\]\((.+?)\)", stripped)
        if image_match:
            flush()
            blocks.append(("image", (image_match.group(1), (path.parent / image_match.group(2)).resolve())))
            index += 1
            continue
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(("heading", (level, stripped[level:].strip())))
            index += 1
            continue
        if stripped.startswith("|"):
            flush()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            blocks.append(("table", parse_table(table_lines)))
            continue
        bullet = re.match(r"^[-*]\s+(.+)", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)", stripped)
        if bullet or numbered:
            flush()
            item_lines = [bullet.group(1) if bullet else numbered.group(1)]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip() or not re.match(r"^(?: {2,}|\t)\S", continuation):
                    break
                item_lines.append(continuation.strip())
                index += 1
            list_text = " ".join(item_lines)
            blocks.append(("bullet" if bullet else "number", convert_citations(list_text)))
            continue
        paragraph_lines.append(stripped)
        index += 1
    flush()
    return blocks


def insert_chapters(doc: Document, anchor) -> tuple[list[str], list[str]]:
    headings: list[str] = []
    captions: list[str] = []
    table_count: dict[int, int] = {}
    current_chapter = 0
    for chapter_path in sorted((ROOT / "docs" / "dissertation").glob("0[1-6]_*.md")):
        blocks = collect_blocks(chapter_path)
        numbered_list_id = None
        for block_index, (block_type, payload) in enumerate(blocks):
            if block_type != "number":
                numbered_list_id = None
            if block_type == "heading":
                level, text = payload
                if level == 1:
                    match = re.match(r"(\d+)\.\s+(.+)", text)
                    if not match:
                        continue
                    current_chapter = int(match.group(1))
                    text = f"Chapter {current_chapter}: {match.group(2).title()}"
                    paragraph = add_paragraph_before(doc, anchor, text, "Heading 1")
                else:
                    paragraph = add_paragraph_before(doc, anchor, text, f"Heading {min(level, 3)}")
                headings.append(text)
                add_bookmark(doc, paragraph, bookmark_name("heading", text))
                paragraph.paragraph_format.keep_with_next = True
                if text == "4.11 Explanation-release audit":
                    paragraph.paragraph_format.page_break_before = True
            elif block_type == "paragraph":
                text = payload
                if re.match(r"\*\*Figure\s+\d+\.\d+\.\*\*", text):
                    clean = text.replace("**", "")
                    paragraph = add_paragraph_before(doc, anchor, style="DMU Caption")
                    apply_inline(paragraph, clean)
                    paragraph.paragraph_format.keep_together = True
                    figure_match = re.match(r"(Figure \d+\.\d+)\.", clean)
                    if figure_match is None:
                        raise ValueError(f"invalid figure caption: {clean}")
                    figure_number = figure_match.group(1)
                    add_bookmark(doc, paragraph, bookmark_name("figure", figure_number))
                    captions.append(figure_number)
                else:
                    paragraph = add_paragraph_before(doc, anchor, text)
                    paragraph.paragraph_format.widow_control = True
                    if (
                        block_index + 1 < len(blocks)
                        and blocks[block_index + 1][0] in {"bullet", "number"}
                    ):
                        paragraph.paragraph_format.keep_with_next = True
            elif block_type in {"bullet", "number"}:
                style = "DMU Bullet" if block_type == "bullet" else "DMU Number"
                paragraph = add_paragraph_before(doc, anchor, payload, style)
                if block_type == "number":
                    if numbered_list_id is None:
                        numbered_list_id = add_numbering_definition(doc, bullet=False)
                    set_paragraph_numbering(paragraph, numbered_list_id)
                paragraph.paragraph_format.keep_together = True
                next_is_same_list = (
                    block_index + 1 < len(blocks)
                    and blocks[block_index + 1][0] == block_type
                )
                if current_chapter == 5 and block_type == "number" and next_is_same_list:
                    paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_after = Pt(2 if next_is_same_list else 6)
            elif block_type == "code":
                if not add_equation_before(doc, anchor, payload):
                    paragraph = add_paragraph_before(doc, anchor, payload)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.left_indent = Inches(0.35)
                    paragraph.paragraph_format.right_indent = Inches(0.35)
                    paragraph.paragraph_format.keep_together = True
                    for run in paragraph.runs:
                        run.font.name = "Courier New"
                        run.font.size = Pt(9)
            elif block_type == "image":
                alt_text, image_path = payload
                add_image_before(doc, anchor, image_path, alt_text)
            elif block_type == "table":
                table_count[current_chapter] = table_count.get(current_chapter, 0) + 1
                number = table_count[current_chapter]
                title = TABLE_TITLES[(current_chapter, number)]
                caption = f"Table {current_chapter}.{number}: {title}"
                table_number = f"Table {current_chapter}.{number}"
                caption_paragraph = add_table_before(doc, anchor, payload, caption)
                add_bookmark(doc, caption_paragraph, bookmark_name("table", table_number))
                captions.append(table_number)
    return headings, captions


def bookmark_name(kind: str, label: str) -> str:
    """Return a deterministic Word-safe bookmark name."""
    clean = re.sub(r"[^A-Za-z0-9_]", "_", label)
    clean = re.sub(r"_+", "_", clean).strip("_")
    return f"nav_{kind}_{clean}"[:40]


def add_bookmark(doc: Document, paragraph: Paragraph, name: str) -> None:
    """Attach a bookmark spanning one paragraph."""
    bookmark_ids = []
    for element in doc.element.body.iter(qn("w:bookmarkStart")):
        value = element.get(qn("w:id"))
        if value is not None and value.isdigit():
            bookmark_ids.append(int(value))
    bookmark_id = str(max(bookmark_ids, default=0) + 1)

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bookmark_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)

    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph: Paragraph, text: str, target: str) -> None:
    """Append an internal Word hyperlink without changing thesis typography."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), target)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    no_proof = OxmlElement("w:noProof")
    run_properties.append(no_proof)
    run.append(run_properties)

    label, page = text.rsplit("\t", 1)
    label_text = OxmlElement("w:t")
    label_text.set(qn("xml:space"), "preserve")
    label_text.text = label
    run.append(label_text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph.add_run(f"\t{page}")


def set_outline_level(paragraph: Paragraph, level: int = 0) -> None:
    """Expose a front-matter title in Word/PDF navigation panes."""
    p_pr = paragraph._p.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def add_front_entry(
    doc: Document,
    anchor,
    label: str,
    page: str,
    indent: float = 0.0,
    target: str | None = None,
):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(indent)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    entry_text = f"{label}\t{page}"
    if target:
        add_internal_hyperlink(paragraph, entry_text, target)
    else:
        paragraph.add_run(entry_text)
    move_before(paragraph._p, anchor)
    return paragraph


def replace_front_lists(doc: Document, headings: list[str]) -> None:
    _, list_tables = clear_between(doc, "LIST OF FIGURES", "LIST OF TABLES")
    for number, title in FIGURES:
        add_front_entry(
            doc,
            list_tables,
            f"{number}  {title}",
            FIGURE_PAGES.get(number, "?"),
            target=bookmark_name("figure", number),
        )

    list_tables_heading = next(
        p for p in doc.paragraphs if p.text == "LIST OF TABLES"
    )
    list_tables_heading.paragraph_format.page_break_before = True
    _, list_abbr = clear_between(doc, "LIST OF TABLES", "LIST OF ABBREVIATIONS")
    for number, title in TABLES:
        add_front_entry(
            doc,
            list_abbr,
            f"{number}  {title}",
            TABLE_PAGES.get(number, "?"),
            target=bookmark_name("table", number),
        )

    _, abstract_anchor = clear_between(doc, "LIST OF ABBREVIATIONS", "ABSTRACT")
    abbreviations_heading = next(
        p for p in doc.paragraphs if p.text == "LIST OF ABBREVIATIONS"
    )
    abbreviations_heading.paragraph_format.page_break_before = True
    abbreviation_rows = [
        ["Abbreviation", "Full form"],
        ["AoI", "Age of Information"],
        ["BMG-Q", "Localized Bipartite Match Graph Attention Q-learning"],
        ["CI", "Confidence Interval"],
        ["CTDE", "Centralized Training with Decentralized Execution"],
        ["DEF", "Decision-level Explanation Faithfulness"],
        ["GAT", "Graph Attention Network"],
        ["GNN", "Graph Neural Network"],
        ["GxI", "Gradient x Input"],
        ["LOO", "Leave-One-Out"],
        ["MADDPG", "Multi-Agent Deep Deterministic Policy Gradient"],
        ["MAPPO", "Multi-Agent Proximal Policy Optimization"],
        ["MARL", "Multi-Agent Reinforcement Learning"],
        ["MLP", "Multilayer Perceptron"],
        ["PPO", "Proximal Policy Optimization"],
        ["RL", "Reinforcement Learning"],
        ["SUMO", "Simulation of Urban Mobility"],
        ["WAMSN", "Weighted Attention Mass on Stale Nodes"],
        ["XRL", "Explainable Reinforcement Learning"],
    ]
    table = doc.add_table(rows=len(abbreviation_rows), cols=2)
    table.style = None
    set_table_geometry(table, [1900, 7120])
    for row_index, values in enumerate(abbreviation_rows):
        row = table.rows[row_index]
        keep_row(row)
        if row_index == 0:
            repeat_header(row)
        for cell, value in zip(row.cells, values):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            run = cell.paragraphs[0].add_run(value)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            if row_index == 0:
                run.bold = True
                set_cell_shading(cell, "E7E6E6")
    move_before(table._tbl, abstract_anchor)

    abstract_heading = next(p for p in doc.paragraphs if p.text == "ABSTRACT")
    abstract_heading.paragraph_format.page_break_before = True

    toc = find_body_child(doc, "TABLE OF CONTENTS")
    toc_heading = next(p for p in doc.paragraphs if p.text == "TABLE OF CONTENTS")
    toc_heading.paragraph_format.page_break_before = True
    chapter_one = find_body_child(doc, "Chapter 1: Introduction")
    between = list(doc.element.body.iterchildren())
    start_index = between.index(toc)
    end_index = between.index(chapter_one)
    section_properties = None
    for child in between[start_index + 1:end_index]:
        sect_pr = child.find(qn("w:pPr") + "/" + qn("w:sectPr")) if child.tag == qn("w:p") else None
        if sect_pr is not None:
            section_properties = deepcopy(sect_pr)
        doc.element.body.remove(child)
    for heading in headings:
        if heading.startswith("Chapter "):
            add_front_entry(
                doc,
                chapter_one,
                heading,
                HEADING_PAGES.get(heading, "?"),
                target=bookmark_name("heading", heading),
            )
        else:
            add_front_entry(
                doc,
                chapter_one,
                heading,
                HEADING_PAGES.get(heading, "?"),
                indent=0.25,
                target=bookmark_name("heading", heading),
            )
    for title in ("List of Publications", "References", "Appendices"):
        add_front_entry(doc, chapter_one, title, HEADING_PAGES.get(title, "?"),
                        target=bookmark_name("heading", title))
        destination = next(p for p in doc.paragraphs if p.text == title)
        add_bookmark(doc, destination, bookmark_name("heading", title))
    if section_properties is None:
        raise ValueError("front-matter section break not found")
    section_break = doc.add_paragraph()
    section_break._p.get_or_add_pPr().append(section_properties)
    move_before(section_break._p, chapter_one)

    for title, kind in (
        ("TABLE OF CONTENTS", "toc"),
        ("LIST OF FIGURES", "list_of_figures"),
        ("LIST OF TABLES", "list_of_tables"),
    ):
        paragraph = next(p for p in doc.paragraphs if p.text == title)
        add_bookmark(doc, paragraph, bookmark_name("front", kind))
        set_outline_level(paragraph)


def replace_abstract(doc: Document) -> None:
    _, toc = clear_between(doc, "ABSTRACT", "TABLE OF CONTENTS")
    for block_type, payload in collect_blocks(ROOT / "docs" / "dissertation" / "00_abstract.md"):
        if block_type == "heading":
            continue
        if block_type == "paragraph":
            paragraph = add_paragraph_before(doc, toc, payload)
            paragraph.paragraph_format.widow_control = True
    keywords = (
        "Keywords: explainable reinforcement learning; graph attention; "
        "multi-agent reinforcement learning; fleet dispatch; telemetry "
        "degradation; faithfulness."
    )
    paragraph = add_paragraph_before(doc, toc)
    label, terms = keywords.split(":", 1)
    paragraph.add_run(label + ":").bold = True
    paragraph.add_run(terms)


def update_front_text(doc: Document) -> None:
    replacements = {
        "Supervisor: Farhan S. Ujager": "Supervisor: Dr Farhan S. Ujager",
        "submitted in partial fulfilment of the requirements for the award of the degree of":
            "submitted in partial fulfillment of the requirements for the award of the degree of",
        "I declare that this dissertation is my own original work carried out under the supervision of Farhan S. Ujager and that all sources and references have been acknowledged appropriately. This dissertation has not been submitted previously for academic credit at De Montfort University, Dubai, or any other institution.":
            "I declare that this dissertation is my own original work carried out under the supervision of Dr Farhan S. Ujager and that all sources and references have been acknowledged appropriately. This dissertation has not been submitted previously for academic credit at De Montfort University, Dubai, or any other institution.",
        "Master of Science (MSc)": "MSc in Artificial Intelligence",
        "No publication is claimed in this dissertation draft.":
            "No publication is claimed in this dissertation.",
        "ACKNOWLEDGEMENT": "ACKNOWLEDGMENT",
        "August 2026": "September 2026",
    }
    for paragraph in doc.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]
    acknowledgment = next(p for p in doc.paragraphs if p.text == "ACKNOWLEDGMENT")
    body = acknowledgment._p.getnext()
    while body is not None and body.tag != qn("w:p"):
        body = body.getnext()
    if body is not None:
        for child in list(body):
            if child.tag == qn("w:r"):
                body.remove(child)
        paragraph = next(p for p in doc.paragraphs if p._p is body)
        paragraph.add_run(
            "I would like to express my sincere gratitude to my supervisor, "
            "Dr Farhan S. Ujager, for his valuable support and guidance "
            "throughout this research. His advice, encouragement and careful "
            "feedback helped me refine the experiment, interpret the findings "
            "responsibly and complete this dissertation."
        )


def style_references_and_appendices(doc: Document) -> None:
    references = False
    appendices = False
    for paragraph in doc.paragraphs:
        if paragraph.text == "References":
            references = True
            continue
        if paragraph.text == "Appendices":
            references = False
            appendices = True
            continue
        if references and paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(2)
        if appendices and paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if paragraph.text == "Appendix B: Supporting Artefacts":
            paragraph.text = "Appendix B: Supporting Artifacts"
        if paragraph.text.startswith((
            "The complete reproduction procedure", "The complete procedure"
        )):
            paragraph.text = (
                "The complete procedure is provided in docs/REPRODUCE_EXPERIMENTS.md. "
                "The guide identifies the required environment, fixed configurations, "
                "commands, validation checks, and expected outputs. The final audit "
                "uses configs/experiments/dissertation_v12_full_rerun.toml, and its "
                "version-controlled results are stored under "
                "results/dissertation_v12_full_rerun/."
            )
        if paragraph.text.startswith((
            "The project repository contains", "The repository contains"
        )):
            paragraph.text = (
                "The repository contains the experiment configuration, source "
                "code, validation selections, preflight reports, per-seed "
                "analyses, summary tables, figure sources, selected checkpoints, "
                "and archived raw cells used in this dissertation. The release "
                "package includes a SHA-256 manifest for verification."
            )


def replace_references(doc: Document) -> None:
    """Replace the template's inherited bibliography with cited sources."""
    _, appendices = clear_between(doc, "References", "Appendices")
    source = ROOT / "docs" / "dissertation" / "references.md"
    for block_type, payload in collect_blocks(source):
        if block_type != "paragraph":
            continue
        paragraph = add_paragraph_before(doc, appendices, payload)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(2)


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def ensure_image_alt_text(doc: Document) -> None:
    for index, shape in enumerate(doc.inline_shapes):
        doc_pr = shape._inline.docPr
        if not doc_pr.get("descr") and not doc_pr.get("title"):
            description = (
                "De Montfort University Dubai logo"
                if index == 0
                else "Dissertation figure"
            )
            doc_pr.set("descr", description)


def build(base: Path, output: Path) -> None:
    page_map = ROOT / "docs" / "dissertation" / "rendered_page_map.json"
    if page_map.exists():
        maps = json.loads(page_map.read_text())
        FIGURE_PAGES.update(maps["figures"])
        TABLE_PAGES.update(maps["tables"])
        HEADING_PAGES.update(maps["headings"])
    doc = Document(base)
    style_document(doc)
    update_front_text(doc)
    replace_abstract(doc)

    anchor = remove_range_inclusive(doc, "Chapter 1: Introduction", "List of Publications")
    headings, _ = insert_chapters(doc, anchor)
    replace_front_lists(doc, headings)
    replace_references(doc)
    style_references_and_appendices(doc)
    add_statistical_appendix(doc)
    enable_field_updates(doc)
    ensure_image_alt_text(doc)

    doc.core_properties.title = (
        "When Explanations Outlive Their Data: Faithfulness Decoupling in "
        "Graph-Attention MARL Fleet Dispatch under Telemetry Degradation"
    )
    doc.core_properties.author = "Hongwei Lin"
    doc.core_properties.subject = "MSc Dissertation, De Montfort University Dubai"
    doc.core_properties.keywords = "GAT-MARL, faithfulness, telemetry degradation, AoI"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


def add_statistical_appendix(doc: Document) -> None:
    anchor = doc.element.body.sectPr
    search_note = ROOT / "docs/dissertation/literature_search_note.md"
    for kind, payload in collect_blocks(search_note):
        if kind == "heading":
            add_paragraph_before(doc, anchor, payload[1], "Heading 3")
        elif kind == "paragraph":
            add_paragraph_before(doc, anchor, payload)
    paragraph = add_paragraph_before(doc, anchor, "Appendix C: Statistical details", "Heading 2")
    paragraph.paragraph_format.page_break_before = True
    add_paragraph_before(doc, anchor,
        "Table C.1 reproduces the rerun checkpoint-specific H1-H4 statistics. "
        "H1 and H3 report Spearman rho; H2 reports the clean-relative rate difference "
        "and remains exploratory; H4 reports mean within-episode rho. N is the number "
        "of scored records for H1/H3, comparison cells for H2, and episodes for H4. "
        "B is the episode-block count where available. Holm adjustment is within "
        "each checkpoint's H1-H4 family. These are not cross-training-seed tests.")
    source = ROOT / "results/dissertation_v12_full_rerun/supplementary_review/hypothesis_statistics.csv"
    rows = [["Model/seed", "H", "Statistic", "Raw p", "Holm p", "N / B"]]
    for row in csv.DictReader(source.open()):
        label = ("GAT-O" if row["checkpoint"].startswith("H5") else "GAT") + "/" + row["checkpoint"][-2:]
        rows.append([label, row["hypothesis"], f'{float(row["statistic"]):+.4g}',
                     f'{float(row["raw_p"]):.4g}', f'{float(row["holm_p"]):.4g}',
                     row["n"] + " / " + (row["n_episode_blocks"] or "n/a")])
    caption = add_table_before(doc, anchor, rows, "Table C.1: Rerun per-checkpoint hypothesis statistics")
    add_bookmark(doc, caption, bookmark_name("table", "Table C.1"))
    add_paragraph_before(doc, anchor,
        "GAT-O denotes GAT-Outage. Full-precision values and the H1/H3 "
        "confidence intervals are retained in hypothesis_statistics.csv under "
        "results/dissertation_v12_full_rerun/supplementary_review/. Missing "
        "intervals are not estimated retrospectively. The same directory contains "
        "all probability and margin no-op estimates, sample counts and input hashes. "
        "The completed five-seed, 50-epoch protocol and runtime are documented in "
        "docs/FULL_RERUN_V12.md.")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.base, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
